#!/usr/bin/env python3
"""Export public-safe portfolio.json from Christopher's 10% Plan workbook.

Privacy (HARD): never write dollar totals, market values, share counts,
cost basis, $ gains, or absolute portfolio values. Only symbol, sector,
weight %, last price, day change %, aggregate weighted day %, and %
performance series / horizons. Never publish the word "Insurance" —
sleeve labeled Other / Alternative investments.

Performance method (NOT Plot):
  Weight-snapshot backtest = current 10% Plan weights × public daily
  *adjusted* close (Yahoo chart API adjclose = total return incl.
  dividends/distributions). Labeled as reconstruction, not actual book /
  trade history. Cash sleeve = USFR ETF total return. Alternative
  investments (Other) = 4.86% annual interest accrued in total-return
  horizons/chart (but not fabricated into the displayed day %). NLY sleeve displays
  as REITs (Yahoo=NLY).
  Benchmark = SPY adj close (UI label SPX).

Source of truth: HIS portfolio workbook (10% Plan.xlsx), morning-synced from
CPMOffice Daily Data when numbers change.
  Default: /workspace/uploads/10pct-plan.xlsx

Usage:
  python3 scripts/export_portfolio.py [xlsx_path] [out_json]
"""

from __future__ import annotations

import json
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl required: pip install openpyxl")

CT = ZoneInfo("America/Chicago")
UA = "Mozilla/5.0 (compatible; OzarkWireExport/1.0)"

# Plan symbol → Yahoo chart ticker
YAHOO_TICKERS = {
    "BTC": "BTC-USD",
    "Gold": "GC=F",
    "Silver": "SI=F",
    "MUEL": "MUEL",
    "CASH": "USFR",  # cash sleeve = USFR ETF total return
    "REITs": "NLY",  # display label; price history still NLY
}

# Plan symbol → Sheet quote ticker for Change %
SHEET_ALIASES = {
    "BTC": "BTC-USD",
    "Gold": "GC=F",
    "Silver": "SI=F",
}

SKIP_SYMBOLS = {
    "total",
    "energy total",
    "energy",
}

# Alternative investments (never labeled Insurance in export). Interest is accrued
# in total-return horizons/chart only; the displayed sleeve day % remains 0.
FLAT_SYMBOLS = {
    "OTHER",
}
ALT_INTEREST_ANNUAL = 0.0486

ALT_DISPLAY = "Other"  # UI/export label; sector = Alternative investments
CASH_YAHOO = "USFR"
# SPY adj close ≈ S&P 500 total return (dividends in). UI still says SPX.
SPX_YAHOO = "SPY"
MUEL_SYMBOL = "MUEL"
DISPLAY_ALIASES = {
    "NLY": "REITs",  # holdings UI label; Yahoo ticker remains NLY
}


def _is_skip(symbol: str) -> bool:
    s = (symbol or "").strip().lower()
    if not s:
        return True
    if s in SKIP_SYMBOLS:
        return True
    if s.startswith("energy total") or s.startswith("health"):
        return True
    if "gains" in s or "total market" in s:
        return True
    return False


def _to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s or s in ("#REF!", "#N/A", "—", "-", " "):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _cell_str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def load_sheet_quotes(ws) -> dict:
    """Sheet tab: col D ticker, E price, G Change % (fraction)."""
    quotes = {}
    for r in range(2, (ws.max_row or 0) + 1):
        ticker = _cell_str(ws.cell(r, 4).value)
        if not ticker:
            continue
        quotes[ticker.upper()] = {
            "price": _to_float(ws.cell(r, 5).value),
            "dayFrac": _to_float(ws.cell(r, 7).value),
        }
    return quotes


def load_holdings_raw(ws, quotes: dict) -> list:
    """Load holdings with internal mktValue for weight rebase only.
    Includes Alternative investments sleeve as Other. Stops before Total.
    mktValue NEVER exported. Never emits symbol 'Insurance'.
    """
    holdings = []
    for r in range(3, (ws.max_row or 0) + 1):
        symbol = _cell_str(ws.cell(r, 2).value)
        if not symbol:
            continue
        if symbol.lower() == "total":
            break
        if _is_skip(symbol):
            continue

        sector = _cell_str(ws.cell(r, 3).value)
        mkt = _to_float(ws.cell(r, 5).value)
        price = _to_float(ws.cell(r, 4).value)

        if symbol.lower().startswith("cash"):
            display = "CASH"
            sector = "Cash"
        elif symbol.lower() == "insurance":
            display = ALT_DISPLAY
            sector = "Alternative investments"
        elif symbol.upper() == "NLY":
            display = DISPLAY_ALIASES["NLY"]  # REITs in UI; Yahoo still NLY
            sector = sector or "Mortgage REIT"
        elif symbol in ("Gold", "Silver", "BTC"):
            display = symbol.upper() if symbol == "BTC" else symbol
        else:
            display = symbol.strip()

        quote_key = SHEET_ALIASES.get(symbol, symbol)
        q = (
            quotes.get(quote_key.upper())
            or quotes.get(symbol.upper())
            or quotes.get(display.upper())
            or (quotes.get("NLY") if display == "REITs" else None)
        )
        # Cash day% from USFR if present on Sheet
        if display == "CASH":
            q = quotes.get("USFR") or q
        day_frac = q["dayFrac"] if q else None
        if price is None and q and q.get("price") is not None:
            price = q["price"]

        # Flat marks: Alternative investments → 0% day, no public last price
        if display.upper() in FLAT_SYMBOLS or display == ALT_DISPLAY:
            day_frac = 0.0
            price = None

        if display == "CASH":
            # price shown as —; day% filled from USFR after Yahoo fetch if needed
            price = None

        if mkt is None or mkt < 0:
            continue

        holdings.append(
            {
                "symbol": display,
                "sector": sector or "",
                "_mkt": mkt,  # stripped before export
                "price": None if price is None else round(price, 4),
                "dayPct": None if day_frac is None else round(day_frac * 100, 4),
            }
        )
    return holdings


def assign_weights(holdings: list) -> list:
    """Rebase weights onto full book incl. Alternative investments. Drop _mkt."""
    total = sum(h["_mkt"] for h in holdings)
    if total <= 0:
        raise SystemExit("holdings market-value sum is zero — cannot weight")
    out = []
    for h in holdings:
        row = {
            "symbol": h["symbol"],
            "sector": h["sector"],
            "weight": round(h["_mkt"] / total * 100.0, 4),
            "price": h["price"],
            "dayPct": h["dayPct"],
        }
        out.append(row)
    return out


def day_weighted_pct(holdings: list, exclude: set[str] | None = None) -> float | None:
    excl = {s.upper() for s in (exclude or set())}
    with_day = [
        h
        for h in holdings
        if h.get("dayPct") is not None and h.get("symbol", "").upper() not in excl
    ]
    if not with_day:
        return None
    wsum = sum(h["weight"] for h in with_day)
    if wsum <= 0:
        return None
    return round(sum(h["weight"] / wsum * h["dayPct"] for h in with_day), 4)


def yahoo_ticker_for(symbol: str) -> str | None:
    if symbol.upper() in FLAT_SYMBOLS or symbol == ALT_DISPLAY:
        return None  # synthetic flat
    if symbol.upper() == "CASH":
        return CASH_YAHOO
    return YAHOO_TICKERS.get(symbol, symbol)


def fetch_yahoo_daily(ticker: str, range_str: str = "5y") -> dict[str, float]:
    """Return {ISO date: adjclose} (total return). Fall back to close. Empty on failure."""
    url = (
        "https://query1.finance.yahoo.com/v8/finance/chart/"
        + urllib.parse.quote(ticker, safe="=^")
        + f"?interval=1d&range={range_str}"
    )
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=45) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, json.JSONDecodeError) as e:
        print(f"  warn: yahoo fetch failed for {ticker}: {e}", file=sys.stderr)
        return {}

    result = (payload.get("chart") or {}).get("result") or []
    if not result:
        err = (payload.get("chart") or {}).get("error")
        print(f"  warn: yahoo no result for {ticker}: {err}", file=sys.stderr)
        return {}
    r0 = result[0]
    ts = r0.get("timestamp") or []
    ind = r0.get("indicators") or {}
    adj = ((ind.get("adjclose") or [{}])[0] or {}).get("adjclose")
    closes = ((ind.get("quote") or [{}])[0] or {}).get("close") or []
    # Prefer adjclose (dividends/distributions reinvested) when any non-null pts exist
    adj_ok = adj and any(v is not None for v in adj)
    series = adj if adj_ok else closes
    out: dict[str, float] = {}
    for t, c in zip(ts, series):
        if c is None:
            continue
        d = datetime.fromtimestamp(t, tz=timezone.utc).date().isoformat()
        out[d] = float(c)
    return out


def last_day_pct(series: dict[str, float]) -> float | None:
    """Trailing 1-session % from a date→price map."""
    if len(series) < 2:
        return None
    keys = sorted(series.keys())
    p0, p1 = series[keys[-2]], series[keys[-1]]
    if not p0:
        return None
    return round((p1 / p0 - 1.0) * 100.0, 4)


def forward_fill(dates: list[str], series: dict[str, float]) -> list[float | None]:
    vals: list[float | None] = []
    last = None
    for d in dates:
        if d in series:
            last = series[d]
        vals.append(last)
    return vals


def build_calendar(price_maps: dict[str, dict[str, float]], start: date, end: date) -> list[str]:
    """Union of all trading dates in [start, end], sorted."""
    s = start.isoformat()
    e = end.isoformat()
    dates = set()
    for m in price_maps.values():
        for d in m:
            if s <= d <= e:
                dates.add(d)
    return sorted(dates)


def portfolio_daily_returns(
    holdings: list,
    price_maps: dict[str, dict[str, float]],
    dates: list[str],
    exclude: set[str] | None = None,
) -> list[float | None]:
    """Weight-snapshot returns; accrue Alternatives interest in horizons/chart.

    The holdings table still reports the supplied Alternatives day % (0.0);
    the 4.86% annual income assumption belongs only in total-return series.
    """
    excl = {s.upper() for s in (exclude or set())}
    active = [h for h in holdings if h["symbol"].upper() not in excl]
    wsum = sum(h["weight"] for h in active)
    if wsum <= 0 or len(dates) < 2:
        return [None] * len(dates)

    paths: dict[str, list[float | None]] = {}
    for h in active:
        sym = h["symbol"]
        if sym.upper() in FLAT_SYMBOLS or sym == ALT_DISPLAY:
            paths[sym] = [1.0] * len(dates)  # flat
            continue
        ysym = yahoo_ticker_for(sym)
        series = price_maps.get(ysym or "", {})
        paths[sym] = forward_fill(dates, series)

    rets: list[float | None] = [None]  # first day no return
    for i in range(1, len(dates)):
        day_r = 0.0
        used_w = 0.0
        for h in active:
            sym = h["symbol"]
            w = h["weight"] / wsum
            path = paths[sym]
            p0, p1 = path[i - 1], path[i]
            if p0 is None or p1 is None or p0 == 0:
                continue
            if sym.upper() == ALT_DISPLAY.upper():
                # Christopher reports ~4.86% interest for this sleeve. Accrue it
                # over the actual date gap so weekends/non-trading gaps count,
                # without inventing a displayed day % for the holding.
                gap_days = max(0, (date.fromisoformat(dates[i]) - date.fromisoformat(dates[i - 1])).days)
                r = (1.0 + ALT_INTEREST_ANNUAL) ** (gap_days / 365.0) - 1.0
            elif sym.upper() in FLAT_SYMBOLS:
                r = 0.0
            else:
                r = p1 / p0 - 1.0
            day_r += w * r
            used_w += w
        if used_w <= 0:
            rets.append(None)
        else:
            rets.append(day_r / used_w)
    return rets


def series_daily_returns(
    price_map: dict[str, float], dates: list[str]
) -> list[float | None]:
    """Daily returns for a single price series (e.g. SPX), LOCF."""
    path = forward_fill(dates, price_map)
    rets: list[float | None] = [None]
    for i in range(1, len(dates)):
        p0, p1 = path[i - 1], path[i]
        if p0 is None or p1 is None or p0 == 0:
            rets.append(None)
        else:
            rets.append(p1 / p0 - 1.0)
    return rets


def cum_series_from(dates: list[str], rets: list[float | None], start_date: str) -> list:
    """Cumulative % series from first date >= start_date. Index 0 = 0%."""
    idx = None
    for i, d in enumerate(dates):
        if d >= start_date:
            idx = i
            break
    if idx is None:
        return []
    cum = 1.0
    out = [{"date": dates[idx], "pct": 0.0}]
    for i in range(idx + 1, len(dates)):
        r = rets[i]
        if r is None:
            continue
        cum *= 1.0 + r
        out.append({"date": dates[i], "pct": round((cum - 1.0) * 100.0, 4)})
    return out


def total_return_pct(dates: list[str], rets: list[float | None], start_date: str) -> float | None:
    series = cum_series_from(dates, rets, start_date)
    if len(series) < 2:
        if len(series) == 1:
            return 0.0
        return None
    return series[-1]["pct"]


def earliest_data_date(dates: list[str], rets: list[float | None]) -> str | None:
    for i in range(1, len(dates)):
        if rets[i] is not None:
            return dates[i - 1] if i > 0 else dates[i]
    return dates[0] if dates else None


def horizon_row(
    label: str,
    dates: list[str],
    rets: list[float | None],
    rets_ex: list[float | None],
    rets_spx: list[float | None],
    start_date: str,
    need_by: str | None,
) -> dict:
    """need_by: if earliest price data is after this ISO date, mark unavailable."""
    earliest = earliest_data_date(dates, rets)
    out = {
        "label": label,
        "pct": None,
        "pctExMuel": None,
        "spxPct": None,
        "vsSpx": None,
        "vsSpxExMuel": None,
        "status": "unavailable",
        "reason": None,
    }
    if not earliest:
        out["reason"] = "n/a — no public price history"
        return out
    if need_by and earliest > need_by:
        out["reason"] = (
            f"n/a — price history starts {earliest}; "
            f"{label} needs data on/before {need_by}"
        )
        return out
    pct = total_return_pct(dates, rets, start_date)
    pct_ex = total_return_pct(dates, rets_ex, start_date)
    spx = total_return_pct(dates, rets_spx, start_date)
    if pct is None:
        out["reason"] = f"n/a — insufficient overlap for {label}"
        return out
    out["pct"] = pct
    out["pctExMuel"] = pct_ex
    out["spxPct"] = spx
    if spx is not None:
        out["vsSpx"] = round(pct - spx, 4)
        if pct_ex is not None:
            out["vsSpxExMuel"] = round(pct_ex - spx, 4)
    out["status"] = "ok"
    out["reason"] = (
        f"current weights × public prices from {max(start_date, earliest)}"
    )
    return out


def export(xlsx: Path, out: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    plan = wb["10% Plan"]
    quotes = load_sheet_quotes(wb["Sheet"])
    raw = load_holdings_raw(plan, quotes)
    holdings = assign_weights(raw)

    print(f"holdings={len(holdings)} (incl. {ALT_DISPLAY})", file=sys.stderr)
    for h in holdings:
        if h["symbol"] in (ALT_DISPLAY, "MUEL", "CASH"):
            print(f"  {h['symbol']}: weight={h['weight']}% day={h['dayPct']}", file=sys.stderr)

    # Fetch public prices (holdings + USFR + SPX)
    tickers = sorted(
        {
            yahoo_ticker_for(h["symbol"])
            for h in holdings
            if yahoo_ticker_for(h["symbol"])
        }
        | {CASH_YAHOO, SPX_YAHOO}
    )
    price_maps: dict[str, dict[str, float]] = {}
    for i, t in enumerate(tickers):
        print(f"fetch [{i+1}/{len(tickers)}] {t}", file=sys.stderr)
        price_maps[t] = fetch_yahoo_daily(t, "5y")
        print(f"  → {len(price_maps[t])} pts", file=sys.stderr)
        time.sleep(0.15)

    # Fill CASH day% from USFR if missing
    usfr_day = last_day_pct(price_maps.get(CASH_YAHOO) or {})
    for h in holdings:
        if h["symbol"] == "CASH" and h.get("dayPct") is None and usfr_day is not None:
            h["dayPct"] = usfr_day

    today = datetime.now(CT).date()
    start = today - timedelta(days=365 * 4 + 30)  # buffer past 3y
    # Include SPX in calendar so benchmark aligns
    dates = build_calendar(price_maps, start, today)
    print(
        f"calendar days={len(dates)} {dates[0] if dates else None} → {dates[-1] if dates else None}",
        file=sys.stderr,
    )

    rets = portfolio_daily_returns(holdings, price_maps, dates)
    rets_ex = portfolio_daily_returns(holdings, price_maps, dates, exclude={MUEL_SYMBOL})
    rets_spx = series_daily_returns(price_maps.get(SPX_YAHOO) or {}, dates)

    ytd_start = date(today.year, 1, 1).isoformat()
    six_m_start = (today - timedelta(days=182)).isoformat()
    one_y_start = (today - timedelta(days=365)).isoformat()
    three_y_start = (today - timedelta(days=365 * 3)).isoformat()

    def _pack_series(start_d: str, label: str) -> dict:
        return {
            "label": label,
            "model": cum_series_from(dates, rets, start_d),
            "exMuel": cum_series_from(dates, rets_ex, start_d),
            "spx": cum_series_from(dates, rets_spx, start_d),
        }

    chart_series = {
        "ytd": _pack_series(ytd_start, f"YTD {today.year}"),
        "sixMonth": _pack_series(six_m_start, "6M"),
        "oneYear": _pack_series(one_y_start, "1Y"),
        "threeYear": _pack_series(three_y_start, "3Y"),
    }

    # Back-compat aliases (YTD) for older UI
    ytd_series = chart_series["ytd"]["model"]
    ytd_series_ex = chart_series["ytd"]["exMuel"]
    ytd_series_spx = chart_series["ytd"]["spx"]

    horizons = {
        "ytd": horizon_row("YTD", dates, rets, rets_ex, rets_spx, ytd_start, None),
        "sixMonth": horizon_row(
            "6M", dates, rets, rets_ex, rets_spx, six_m_start, six_m_start
        ),
        "oneYear": horizon_row(
            "1Y", dates, rets, rets_ex, rets_spx, one_y_start, one_y_start
        ),
        "threeYear": horizon_row(
            "3Y", dates, rets, rets_ex, rets_spx, three_y_start, three_y_start
        ),
    }

    day_w = day_weighted_pct(holdings)
    day_ex = day_weighted_pct(holdings, exclude={MUEL_SYMBOL})

    now = datetime.now(CT)
    missing = [t for t in tickers if not price_maps.get(t)]
    coverage = (
        f"Yahoo daily for {len(tickers) - len(missing)}/{len(tickers)} tickers"
        + (f" (missing: {', '.join(missing)})" if missing else "")
    )

    # Internal/export-only note — UI must not lecture from this field
    _method_note = (
        "Weight-snapshot total-return backtest (export only). Yahoo adjclose; "
        "Cash=USFR; Other (Alternative investments)=4.86% annual interest accrued "
        "in total-return horizons/chart (not displayed day %); NLY shown as REITs; "
        "vs SPY (UI: SPX). Ex-MUEL: drop MUEL, renormalize. No $ totals/shares. "
        + coverage + "."
    )

    payload = {
        "updated": now.isoformat(timespec="seconds"),
        "source": (
            "10% Plan sheet · performance = current weights × Yahoo adj close "
            "(total return). Cash = USFR. Alternatives interest = 4.86% annual. "
            "Benchmark = SPY (labeled SPX). "
            "NLY sleeve labeled REITs."
        ),
        "note": "",  # UI disclaimer removed — privacy stays in code/export only
        "_exportNote": _method_note,
        "method": "weight_snapshot_total_return_adjclose",
        "returnType": "total_return",
        "cashProxy": "USFR",
        "alternativesInterestAnnualPct": 4.86,
        "alternativesInterestMethod": "Accrued in YTD/6M/1Y/3Y total-return chart and horizons; excluded from displayed day %." ,
        "benchmark": "SPY",
        "benchmarkLabel": "SPX",
        "holdings": holdings,
        "dayWeightedPct": day_w,
        "dayWeightedPctExMuel": day_ex,
        "chartSeries": chart_series,
        "chartWindows": ["ytd", "sixMonth", "oneYear", "threeYear"],
        "ytdSeries": ytd_series,
        "ytdSeriesExMuel": ytd_series_ex,
        "ytdSeriesSpx": ytd_series_spx,
        "ytdLatestPct": ytd_series[-1]["pct"] if ytd_series else None,
        "ytdLatestPctExMuel": ytd_series_ex[-1]["pct"] if ytd_series_ex else None,
        "ytdLatestPctSpx": ytd_series_spx[-1]["pct"] if ytd_series_spx else None,
        "chartLabel": f"YTD {today.year}",
        "horizons": horizons,
        "exMuelMethod": (
            "day + horizons + chart: MUEL dropped, remaining weights renormalized; "
            "Other (Alternative investments) accrues 4.86% annual interest in horizons/chart; "
            "Cash = USFR; REITs=NLY"
        ),
    }

    # Strip any accidental private keys (keep _exportNote out of public file too)
    def _scrub(o):
        if isinstance(o, dict):
            return {
                k: _scrub(v)
                for k, v in o.items()
                if not str(k).startswith("_")
                and k
                not in {
                    "mktValue",
                    "marketValue",
                    "shares",
                    "costBasis",
                    "originalCost",
                    "gainLoss",
                    "muelMkt",
                    "plotTotal",
                    "_mkt",
                    "_exportNote",
                }
            }
        if isinstance(o, list):
            return [_scrub(x) for x in o]
        return o

    payload = _scrub(payload)

    text = json.dumps(payload, indent=2) + "\n"
    if "$" in text:
        raise SystemExit("Privacy refuse: '$' found in portfolio.json export")
    for bad in ("mktValue", "shares", "costBasis", "plotTotal", "_mkt"):
        if f'"{bad}"' in text:
            raise SystemExit(f"Privacy refuse: forbidden key {bad}")
    # Never publish the word Insurance in the public JSON
    if "Insurance" in text or "insurance" in text:
        raise SystemExit("Privacy/UI refuse: 'Insurance' found in portfolio.json export")

    def walk(o, path=""):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, f"{path}.{k}")
        elif isinstance(o, list):
            for i, v in enumerate(o):
                walk(v, f"{path}[{i}]")
        elif isinstance(o, (int, float)) and abs(o) > 100_000 and "price" not in path:
            raise SystemExit(f"Privacy refuse: large number at {path}={o}")

    walk(payload)

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    print(f"export note (stderr only): {_method_note}", file=sys.stderr)
    return payload


def main():
    root = Path(__file__).resolve().parents[1]
    xlsx = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path("/workspace/uploads/10pct-plan.xlsx")
    )
    out = (
        Path(sys.argv[2])
        if len(sys.argv) > 2
        else root / "content" / "portfolio.json"
    )
    if not xlsx.exists():
        raise SystemExit(f"xlsx not found: {xlsx}")
    data = export(xlsx, out)
    print(f"Wrote {out}")
    print(
        f"day={data['dayWeightedPct']} dayEx={data['dayWeightedPctExMuel']} "
        f"ytd={data['ytdLatestPct']} ytdEx={data['ytdLatestPctExMuel']} "
        f"ytdSpx={data.get('ytdLatestPctSpx')}"
    )
    for k, row in data["horizons"].items():
        print(
            f"  {k}: status={row['status']} pct={row['pct']} pctEx={row['pctExMuel']} "
            f"spx={row.get('spxPct')} vs={row.get('vsSpx')} vsEx={row.get('vsSpxExMuel')} "
            f"— {row.get('reason')}"
        )


if __name__ == "__main__":
    main()
